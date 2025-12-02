# datalogger.py

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font
import os
from datetime import datetime

# --- NEW: Generate a date-specific filename ---
# The file will be named like "weather_data_2025-08-18.xlsx"
FILENAME = f"weather_data_{datetime.now().strftime('%Y-%m-%d')}.xlsx"

DATA_HEADERS = [
    "Date", "Time", "GPS Coordinates", "Altitude (m)",
    "Relative Humidity (%)", "Globe Temp (C)", "Air Temperature (C)","Adjusted Air Speed (m/s)", "Original Air Speed (m/s)" ,"TMRT (C)", "UTCI"
]
METADATA_LABELS = [
    "SciWalK, Pilot Phase, 2025-2026",
    "Measurement Data",
    "Collection Date",
    "Collector",
    "Measurement Device",
    "Measurement Height",
    "Focus",
    "Weather Conditions",
    "Comments",
]

def setup_workbook(metadata):
    """
    Creates the workbook if it doesn't exist.
    Writes the metadata table and data headers with bold formatting.
    """
    if not os.path.exists(FILENAME):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sensor Data"
        
        # --- Define Bold Font Style ---
        bold_font = Font(bold=True)

        # --- Write Metadata Mini-Table with Bold Labels ---
        for i, label in enumerate(METADATA_LABELS, 1):
            cell = sheet.cell(row=i, column=1, value=label)
            cell.font = bold_font  # Apply bold style

        # Write metadata values in the second column (not bold)
        sheet.cell(row=3, column=2, value=metadata["Collection Date"])
        sheet.cell(row=4, column=2, value=metadata["Collector"])
        sheet.cell(row=5, column=2, value=metadata["Measurement Device"])
        sheet.cell(row=6, column=2, value=metadata["Measurement Height"])
        sheet.cell(row=7, column=2, value=metadata["Focus"])
        sheet.cell(row=8, column=2, value=metadata["Weather Conditions"])
        sheet.cell(row=9, column=2, value=metadata["Comments"])

        # --- Leave blank rows for spacing ---
        sheet.append([]) # Blank row 10
        sheet.append([]) # Blank row 11
        
        # --- Write Data Headers and make them bold ---
        sheet.append(DATA_HEADERS)
        # Get the row number where headers were just added
        header_row = sheet.max_row
        for cell in sheet[header_row]:
            cell.font = bold_font # Apply bold style

        workbook.save(FILENAME)
        print(f"Created new workbook with bold headers: {FILENAME}")
    else:
        print(f"Workbook {FILENAME} already exists. Appending data.")


def append_data(data):
    """Appends a new row of data to the Sensor Data sheet."""
    try:
        workbook = openpyxl.load_workbook(FILENAME)
        sheet = workbook.active
        sheet.append(data)
        workbook.save(FILENAME)
    except FileNotFoundError:
        print(f"Error: {FILENAME} not found.")
    except Exception as e:
        print(f"An error occurred while writing to Excel: {e}")
